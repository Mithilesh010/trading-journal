import io
from datetime import datetime, timedelta
import pytz
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side as BorderSide
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

IST = pytz.timezone('Asia/Kolkata')

def filter_trades_by_range(trades, range_param):
    """Filter trades by range: 15days, 1month, or all."""
    now_ist = datetime.now(IST).date()
    if range_param == '15days':
        start_date = now_ist - timedelta(days=15)
        return [t for t in trades if t.trade_date >= start_date]
    elif range_param == '1month':
        start_date = now_ist - timedelta(days=30)
        return [t for t in trades if t.trade_date >= start_date]
    return trades

def generate_excel_export(trades, user, range_param='all'):
    """Generates a professional Excel workbook for the user's trades."""
    filtered_trades = filter_trades_by_range(trades, range_param)
    # Sort chronologically
    filtered_trades.sort(key=lambda t: (t.trade_date, t.entry_time or '00:00'))

    wb = Workbook()
    ws = wb.active
    ws.title = "Trading Journal"
    ws.views.sheetView[0].showGridLines = True

    # Styling constants
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    green_font = Font(name="Segoe UI", size=10, bold=True, color="16A34A")
    red_font = Font(name="Segoe UI", size=10, bold=True, color="DC2626")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    right_align = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=BorderSide(style='thin', color='CBD5E1'),
        right=BorderSide(style='thin', color='CBD5E1'),
        top=BorderSide(style='thin', color='CBD5E1'),
        bottom=BorderSide(style='thin', color='CBD5E1')
    )

    # Title & Metadata Block
    ws.merge_cells("A1:R1")
    title_cell = ws["A1"]
    title_cell.value = "TRADING TERMINAL - TRADE LOG"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    ws["A2"].value = f"User: {user.name} ({user.email}) | Export Date: {datetime.now(IST).strftime('%d-%b-%Y %I:%M %p')} IST | Range: {range_param.upper()}"
    ws["A2"].font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    ws.merge_cells("A2:R2")
    ws.row_dimensions[2].height = 18

    # Table Headers
    headers = [
        "Date", "Entry Time", "Exit Time", "Instrument", "Side",
        "Quantity", "Entry Price", "Exit Price", "Stop Loss", "Target",
        "P&L", "Risk", "Reward", "R:R", "Strategy", "Setup", "Status", "Notes"
    ]

    header_row = 4
    ws.row_dimensions[header_row].height = 24

    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # Data Rows
    current_row = header_row + 1
    total_pnl = 0.0

    for t in filtered_trades:
        ws.row_dimensions[current_row].height = 20
        pnl = t.pnl
        if pnl is not None:
            total_pnl += pnl

        row_values = [
            t.trade_date.strftime('%d-%m-%Y'),
            t.entry_time or '-',
            t.exit_time or '-',
            t.instrument,
            t.side.upper(),
            t.quantity,
            t.entry_price,
            t.exit_price if t.exit_price is not None else '-',
            t.stop_loss if t.stop_loss is not None else '-',
            t.target if t.target is not None else '-',
            pnl if pnl is not None else '-',
            t.risk if t.risk is not None else '-',
            t.reward if t.reward is not None else '-',
            f"1:{t.rr_ratio}" if t.rr_ratio else '-',
            t.strategy or '-',
            t.setup or '-',
            t.status,
            t.notes or ''
        ]

        for col_idx, val in enumerate(row_values, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = data_font
            cell.border = thin_border

            # Alignment & specialized color styling
            if col_idx in (1, 2, 3, 5, 14, 17):
                cell.alignment = center_align
            elif col_idx in (6, 7, 8, 9, 10, 11, 12, 13):
                cell.alignment = right_align
            else:
                cell.alignment = left_align

            # Color side
            if col_idx == 5:
                cell.font = green_font if t.side.upper() == 'BUY' else red_font
            # Color P&L
            if col_idx == 11 and pnl is not None:
                cell.font = green_font if pnl >= 0 else red_font
                cell.number_format = '#,##0.00'

        current_row += 1

    # Summary Row
    summary_row = current_row + 1
    ws.cell(row=summary_row, column=10, value="Total Realized P&L:").font = bold_font
    ws.cell(row=summary_row, column=10).alignment = right_align
    total_cell = ws.cell(row=summary_row, column=11, value=round(total_pnl, 2))
    total_cell.font = green_font if total_pnl >= 0 else red_font
    total_cell.alignment = right_align
    total_cell.border = thin_border
    total_cell.number_format = '#,##0.00'

    # Auto-adjust column widths
    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for r in range(header_row, current_row + 2):
            cell_val = str(ws.cell(row=r, column=col_idx).value or '')
            max_len = max(max_len, len(cell_val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def generate_pdf_export(trades, user, range_param='all'):
    """Generates a professional Landscape PDF trade report using ReportLab."""
    filtered_trades = filter_trades_by_range(trades, range_param)
    filtered_trades.sort(key=lambda t: (t.trade_date, t.entry_time or '00:00'))

    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(letter),
        rightMargin=24,
        leftMargin=24,
        topMargin=24,
        bottomMargin=24
    )

    elements = []
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A')
    )
    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor('#64748B')
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=1
    )
    cell_left = ParagraphStyle(
        'CellLeft',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        alignment=0
    )

    # Document Header
    elements.append(Paragraph("TRADING TERMINAL - TRADING JOURNAL REPORT", title_style))
    meta_text = (
        f"<b>Trader:</b> {user.name} ({user.email}) &nbsp;|&nbsp; "
        f"<b>Generated:</b> {datetime.now(IST).strftime('%d-%m-%Y %I:%M %p')} IST &nbsp;|&nbsp; "
        f"<b>Range:</b> {range_param.upper()} ({len(filtered_trades)} trades)"
    )
    elements.append(Paragraph(meta_text, meta_style))
    elements.append(Spacer(1, 10))

    # Summary Statistics Bar
    closed = [t for t in filtered_trades if t.status == 'Closed' and t.pnl is not None]
    total_pnl = sum(t.pnl for t in closed)
    wins = len([t for t in closed if t.pnl > 0])
    win_rate = round((wins / len(closed) * 100), 1) if closed else 0.0

    summary_data = [
        ["Total Trades", "Closed Trades", "Win Rate", "Total Realized P&L"],
        [str(len(filtered_trades)), str(len(closed)), f"{win_rate}%", f"{total_pnl:+,.2f}"]
    ]
    summary_table = Table(summary_data, colWidths=[120, 120, 120, 150])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#F1F5F9')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor('#475569')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TEXTCOLOR', (3, 1), (3, 1), colors.HexColor('#16A34A') if total_pnl >= 0 else colors.HexColor('#DC2626')),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(summary_table)
    elements.append(Spacer(1, 12))

    # Trade Table
    table_headers = [
        "Date", "Time", "Instrument", "Side", "Qty",
        "Entry", "Exit", "P&L", "R:R", "Strategy", "Status"
    ]
    table_data = [[Paragraph(f"<b>{h}</b>", cell_style) for h in table_headers]]

    for t in filtered_trades:
        side_color = "#16A34A" if t.side.upper() == 'BUY' else "#DC2626"
        pnl_str = '-'
        pnl_color = "#475569"
        if t.pnl is not None:
            pnl_str = f"{t.pnl:+,.2f}"
            pnl_color = "#16A34A" if t.pnl >= 0 else "#DC2626"

        row = [
            Paragraph(t.trade_date.strftime('%d-%m-%Y'), cell_style),
            Paragraph(t.entry_time or '-', cell_style),
            Paragraph(t.instrument, cell_left),
            Paragraph(f"<font color='{side_color}'><b>{t.side.upper()}</b></font>", cell_style),
            Paragraph(f"{t.quantity:g}", cell_style),
            Paragraph(f"{t.entry_price:g}", cell_style),
            Paragraph(f"{t.exit_price:g}" if t.exit_price is not None else "-", cell_style),
            Paragraph(f"<font color='{pnl_color}'><b>{pnl_str}</b></font>", cell_style),
            Paragraph(f"1:{t.rr_ratio}" if t.rr_ratio else "-", cell_style),
            Paragraph(t.strategy or "-", cell_left),
            Paragraph(f"<b>{t.status}</b>", cell_style)
        ]
        table_data.append(row)

    col_widths = [60, 40, 120, 45, 45, 55, 55, 65, 45, 120, 50]
    trade_table = Table(table_data, colWidths=col_widths, repeatRows=1)
    trade_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
    ]))
    elements.append(trade_table)

    # Footer note with required credit
    elements.append(Spacer(1, 15))
    credit_style = ParagraphStyle(
        'Credit',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor('#94A3B8'),
        alignment=1
    )
    elements.append(Paragraph("Trading Terminal &copy; 2026 &nbsp;|&nbsp; Designed & Developed by Mithilesh Kumar", credit_style))

    doc.build(elements)
    output.seek(0)
    return output
